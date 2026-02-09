from docx import Document
import uuid
import os

def create_doc_from_fields(fields):
    doc = Document()
    doc.add_heading('FORMULIR DIGITAL HASIL OCR', level=1)

    for k, v in fields.items():
        p = doc.add_paragraph()
        p.add_run(f"{k} : ").bold = True
        p.add_run(v)

    filename = f"hasil_form_{uuid.uuid4().hex}.docx"
    path = os.path.join("static", filename)
    doc.save(path)

    return path
